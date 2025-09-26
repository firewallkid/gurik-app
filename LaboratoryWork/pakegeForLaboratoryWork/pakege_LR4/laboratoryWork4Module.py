from math import *
from openpyxl import Workbook
from docx import Document
import sqlite3


def calculate_y(x, b):
    """
    Вычисляет значение заданной функции.
    :param x: Число, для которого нужно вычислить функцию
    :param b: Конец отрезка
    :return: Значение функции y
    """
    if x <= -6 * b:
        return -(x + 3 * b) ** 2 - 2 * b
    else:
        return b * cos(x + 3 * b) - 3 * b


def create_db():
    #Создает базу данных и таблицу для хранения значений x и y
    conn = sqlite3.connect('C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.db')
    cursor = conn.cursor()
    # Создание таблицы, если она не существует
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS calculate_y_table (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            x REAL NOT NULL,
            result REAL NOT NULL
        )
    ''')
    conn.commit()
    conn.close()


def insert_into_db(x, result):
    """
    Вставляет данные в базу данных.
    :param x: Значение x.
    :param result: Значение y.
    """
    conn = sqlite3.connect('C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.db')
    cursor = conn.cursor()
    # Вставка значения x и y
    cursor.execute("INSERT INTO calculate_y_table (x, result) VALUES (?, ?)", (x, result))
    conn.commit()
    conn.close()


def tabulate_y(a, b, h):
    """
    Табулирует заданную функцию и записывает результаты в текстовый файл и Excel.
    :param a: Начало отрезка.
    :param b: Конец отрезка.
    :param h: Шаг.
    """
    total_sum = 0
    total_pr = 1
    total_count = 0
    if h <= 0:
        print('Шаг должен быть положительным числом.')
        return

    # Вывод в консоль
    print(f"{'x':>9} | {'y':>9}")
    print("-" * 23)

    # Сохранение в текстовый файл
    with open("C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.txt", "w") as txt_file:
        txt_file.write(f"{'x':>10} | {'y':>10}\n")
        txt_file.write("-" * 23)

        # Сохранение в Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "calculate_y_table"
        ws.append(["x", "y"])
        # Сохранение в Word
        doc = Document()
        doc.add_heading("Вычисление значения y", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'x'
        hdr_cells[1].text = 'y'
        # Вычисление значений и запись в файлы
        for i in range(int((b - a) / h) + 1):
            x = round(a + i * h, 2)
            result = round(calculate_y(x, b), 4)
            if int(result) > 0:
                total_count += 1
                total_sum += result
                total_pr *= result
            # Вывод в консоль
            print(f"{x:>6f} | {result:>6f}")
            # Запись в текстовый файл
            txt_file.write(f"{x:.2f} | {result:4f}\n")
            # Запись в Excel
            ws.append([x, result])
            # Запись в Word
            row_cells = table.add_row().cells
            row_cells[0].text = f"{x:.4f}"
            row_cells[1].text = f"{result:.4f}"
            # Запись в базу данных
            insert_into_db(x, result)
        # Сохранение файлов
        wb.save("C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.xlsx")
        doc.save("C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.docx")
        conn = sqlite3.connect('C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.db')
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM calculate_y_table")
        results = cursor.fetchall()
        conn.close()
        txt_file.write("\n" + "_" * 23 + "\n")
        txt_file.write(f"Количество положительных y: {total_count:.4f}")
        txt_file.write(f"Сумма значений положительных y: {total_sum:.4f}")
        txt_file.write(f"Произведение значений положительных y: {total_pr:.4f}")

        print(f"Количество положительных y: {total_count:.4f}")
        print(f"Сумма значений положительных y: {total_sum:.4f}")
        print(f"Произведение значений положительных y: {total_pr:.4f}")
        print(
            "\nТаблица сохранена в файлы: \n'calculate_y_table.txt', \n'calculate_y_table.xlsx', \n'calculate_y_table.docx'. \nДанные записаны в таблицу")

def execute_query(query, params=None):
    """
    Это функция для выполнения SQL-запроса.
    В зависимости от того, передаются ли параметры она выполняет запрос с параметрами или без них
    :param query: Строка SQL-запроса.
    :param params: Параметры для SQL-запроса (если есть).
    """
    conn = sqlite3.connect('C:/Users/demak/PycharmProjects/pythonProject/LaboratoryWork/pakegeForLaboratoryWork/pakege_LR4/calculate_y_table.db')
    cursor = conn.cursor()
    # Курсор - это объект, предоставляемый модулем sqlite3.
    # Он отвечает на выполнение SQL-запросов и управление результатами этих запросов.
    # После выполнения запроса, например, SELEcT, курсор содержит результаты запроса,
    # которые можно извлекать методом fetchone()
    if params:
        cursor.execute(query, params)
    else:
        cursor.execute(query)
    # Получаем все строки результата
    rows = cursor.fetchall()
    # Печатаем результаты
    for row in rows:
        print(row)

    conn.close() # Закрываем соединение
def example():
    # Пример запроса 1: Получить все записи из таблицы
    print("Все записи из базы данных:")
    execute_query("SELECT * FROM calculate_y_table;")

    # Пример запроса 2: Получить записи для значений x между 0 и 1
    print("Записи для x между 0 и 1:")
    execute_query("SELECT * FROM calculate_y_table WHERE x BETWEEN ? AND ?", (0, 1))
