from tkinter import *
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime


class ZastavkaFrame(Frame):
    def __init__(self, master, start_callback):
        super().__init__(master)
        self.grid()
        self.start_callback = start_callback
        self.configure_fonts()
        self.create_widgets()

    def configure_fonts(self):
        """Настраивает шрифты для всех элементов"""
        self.main_font = ("Times New Roman", 14)
        self.title_font = ("Times New Roman", 16, "bold")
        self.button_font = ("Times New Roman", 14)

    def create_widgets(self):
        # Основной контейнер для центрирования
        container = Frame(self)
        container.grid(pady=50)

        # Текст заставки
        text = (
            "МИНИСТЕРСТВО ЦИФРОВОГО РАЗВИТИЯ, СВЯЗИ\n"
            "И МАССОВЫХ КОММУНИКАЦИЙ РОССИЙСКОЙ ФЕДЕРАЦИИ\n"
            "Ордена Трудового Красного Знамени федеральное государственное\n"
            "бюджетное образовательное учреждение высшего образования\n"
            "МОСКОВСКИЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ СВЯЗИ И ИНФОРМАТИКИ\n"
        )

        Label(container,
              text=text,
              justify=CENTER,
              font=self.main_font
              ).grid(row=0, column=0, padx=20, pady=10)

        # Заголовок темы
        Label(container,
              text='Тестовая программа на тему "Коммерческий шпионаж"',
              font=self.title_font
              ).grid(row=1, column=0, pady=20)

        # Информация об авторе и проверяющем
        Label(container,
              text="Выполнил: Демаков М. Д.\nПроверил: Гуриков С. Р.",
              font=self.main_font
              ).grid(row=2, column=0, pady=10)

        # Кнопка начала теста
        Button(container,
               text="Начать тест",
               command=self.start_callback,
               font=self.button_font,
               padx=30,
               pady=5
               ).grid(row=3, column=0, pady=30)


class TextBoxFrame(Frame):
    def __init__(self, master, on_next_callback, question_index):
        super().__init__(master)
        self.grid()
        self.on_next_callback = on_next_callback
        self.question_index = question_index
        self.configure_fonts()
        self.create_widgets()

    def configure_fonts(self):
        """Настраивает шрифты для всех элементов"""
        self.main_font = ("Times New Roman", 14)
        self.title_font = ("Times New Roman", 16, "bold")
        self.button_font = ("Times New Roman", 14)

    def check_answer(self):
        """Проверяет правильность ответа пользователя"""
        user_input = self.answer_entry.get().strip().lower()
        correct_answers = [
            "промышленный шпионаж",
            "компьютерный шпионаж",
            "экономический шпионаж",
            "шпионаж"
        ]

        is_correct = any(correct in user_input for correct in correct_answers)

        result = "Вы ответили правильно!" if is_correct else "Вы ответили неправильно!"
        self.result_text.configure(state='normal')
        self.result_text.delete("1.0", END)
        self.result_text.insert("1.0", result)
        self.result_text.configure(state='disabled')
        self.after(1500, lambda: self.load_next_question(is_correct))

    def load_next_question(self, is_correct):
        """Переходит к следующему вопросу"""
        self.destroy()
        self.on_next_callback(is_correct)

    def create_widgets(self):
        """Создает интерфейс для текстового вопроса"""
        # Основной контейнер для центрирования
        container = Frame(self)
        container.grid(pady=20)

        # Заголовок вопроса
        Label(container,
              text=f"Вопрос {self.question_index + 1}:",
              font=self.title_font
              ).grid(row=0, column=0, pady=(0, 20))

        # Текст вопроса
        question_text = (
            "Введите словосочетание, исходя из его определения:\n\n"
            "- противоправное собирание,\n"
            "- похищение, передача, хранение в целях передачи чужой информации,\n"
            "- а также неправомерный доступ к компьютерной информации с целью\n"
            "получения преимуществ для себя или третьих лиц при осуществлении\n"
            "предпринимательской деятельности, нанесения ущерба собственнику,\n"
            "а равно получения материальной или иной выгоды."
        )

        Label(container,
              text=question_text,
              font=self.main_font,
              justify=LEFT
              ).grid(row=1, column=0, pady=(0, 20))

        # Поле для ввода ответа
        self.answer_entry = Entry(container,
                                 width=50,
                                 font=self.main_font)
        self.answer_entry.grid(row=2, column=0, pady=10)

        # Поле для вывода результата
        self.result_text = Text(container,
                               width=50,
                               height=3,
                               font=self.main_font,
                               wrap=WORD,
                               state='disabled')
        self.result_text.grid(row=3, column=0, pady=10)

        # Кнопка проверки
        Button(container,
               text="Ответить",
               command=self.check_answer,
               font=self.button_font,
               padx=30,
               pady=5
               ).grid(row=4, column=0, pady=20)


class CheckQuestionFrame(Frame):
    def __init__(self, master, on_next_callback, question_index):
        super().__init__(master)
        self.grid()
        self.on_next_callback = on_next_callback
        self.question_index = question_index
        self.configure_fonts()
        self.create_widgets()

    def configure_fonts(self):
        """Настраивает шрифты для всех элементов"""
        self.main_font = ("Times New Roman", 14)
        self.title_font = ("Times New Roman", 16, "bold")
        self.button_font = ("Times New Roman", 14)
        self.checkbox_font = ("Times New Roman", 14)

    def proverka(self):
        # Правильные ответы: Дисциплинарная, Материальная, Уголовная, Гражданско-правовая, Административная
        correct = (self.otvet1.get() and  # Дисциплинарная
                   self.otvet2.get() and  # Материальная
                   not self.otvet3.get() and  # Физическая (не должна быть выбрана)
                   self.otvet4.get() and  # Уголовная
                   self.otvet5.get() and  # Гражданско-правовая
                   self.otvet6.get() and  # Административная
                   not self.otvet7.get())  # Трудовая (не должна быть выбрана)

        result = "Вы ответили правильно" if correct else "Вы ответили неправильно"
        self.textbox2.configure(state='normal')
        self.textbox2.delete("1.0", END)
        self.textbox2.insert("1.0", result)
        self.textbox2.configure(state='disabled')
        self.after(1500, lambda: self.load_next(correct))

    def load_next(self, is_correct):
        self.destroy()
        self.on_next_callback(is_correct)

    def create_widgets(self):
        """Создает интерфейс для вопроса с выбором нескольких вариантов"""
        # Основной контейнер для центрирования
        container = Frame(self)
        container.grid(pady=20)

        Label(container,
              text=f"Вопрос {self.question_index + 1}:",
              font=self.title_font
              ).grid(row=0, column=0, pady=(0, 20))

        Label(container,
              text="Правовые основы ответственности делятся на:",
              font=self.main_font
              ).grid(row=1, column=0, pady=(0, 10))

        Label(container,
              text="Выберите несколько вариантов ответа:",
              font=self.main_font
              ).grid(row=2, column=0, pady=(0, 20))

        self.otvet1 = BooleanVar()  # Дисциплинарная
        self.otvet2 = BooleanVar()  # Материальная
        self.otvet3 = BooleanVar()  # Физическая (неверный)
        self.otvet4 = BooleanVar()  # Уголовная
        self.otvet5 = BooleanVar()  # Гражданско-правовая
        self.otvet6 = BooleanVar()  # Административная
        self.otvet7 = BooleanVar()  # Трудовая (неверный)

        Checkbutton(container,
                   text="Дисциплинарная",
                   variable=self.otvet1,
                   font=self.checkbox_font
                   ).grid(row=3, column=0, sticky=W, pady=5)

        Checkbutton(container,
                   text="Материальная",
                   variable=self.otvet2,
                   font=self.checkbox_font
                   ).grid(row=4, column=0, sticky=W, pady=5)

        Checkbutton(container,
                   text="Физическая",
                   variable=self.otvet3,
                   font=self.checkbox_font
                   ).grid(row=5, column=0, sticky=W, pady=5)

        Checkbutton(container,
                   text="Уголовная",
                   variable=self.otvet4,
                   font=self.checkbox_font
                   ).grid(row=6, column=0, sticky=W, pady=5)

        Checkbutton(container,
                   text="Гражданско-правовая",
                   variable=self.otvet5,
                   font=self.checkbox_font
                   ).grid(row=7, column=0, sticky=W, pady=5)

        Checkbutton(container,
                   text="Административная",
                   variable=self.otvet6,
                   font=self.checkbox_font
                   ).grid(row=8, column=0, sticky=W, pady=5)

        Checkbutton(container,
                   text="Трудовая",
                   variable=self.otvet7,
                   font=self.checkbox_font
                   ).grid(row=9, column=0, sticky=W, pady=5)

        self.textbox2 = Text(container,
                            width=50,
                            height=3,
                            font=self.main_font,
                            wrap=WORD,
                            state='disabled')
        self.textbox2.grid(row=10, column=0, pady=20)

        Button(container,
               text="Ответить",
               command=self.proverka,
               font=self.button_font,
               padx=30,
               pady=5
               ).grid(row=11, column=0, pady=10)


class RadioQuestionFrame(Frame):
    def __init__(self, master, on_next_callback, question_index):
        super().__init__(master)
        self.grid()
        self.on_next_callback = on_next_callback
        self.question_index = question_index
        self.configure_fonts()
        self.create_widgets()

    def configure_fonts(self):
        """Настраивает шрифты для всех элементов"""
        self.main_font = ("Times New Roman", 14)
        self.title_font = ("Times New Roman", 16, "bold")
        self.button_font = ("Times New Roman", 14)
        self.radio_font = ("Times New Roman", 14)

    def proverka(self):
        answer = self.otvet.get()
        correct = answer == "Использование надёжных паролей"
        result = "Вы ответили правильно" if correct else "Вы ответили неправильно"
        self.textbox2.configure(state='normal')
        self.textbox2.delete("1.0", END)
        self.textbox2.insert("1.0", result)
        self.textbox2.configure(state='disabled')
        self.after(1500, lambda: self.load_next(correct))

    def load_next(self, is_correct):
        self.destroy()
        self.on_next_callback(is_correct)

    def create_widgets(self):
        """Создает интерфейс для вопроса с выбором одного варианта"""
        # Основной контейнер для центрирования
        container = Frame(self)
        container.grid(pady=20)

        Label(container,
              text=f"Вопрос {self.question_index + 1}:",
              font=self.title_font
              ).grid(row=0, column=0, pady=(0, 20))

        Label(container,
              text="Какой из перечисленных методов является базовым способом защиты ваших данных в социальных сетях?",
              font=self.main_font,
              wraplength=600,
              justify=LEFT
              ).grid(row=1, column=0, pady=(0, 10))

        Label(container,
              text="Выберите правильный ответ:",
              font=self.main_font
              ).grid(row=2, column=0, pady=(0, 20))

        self.otvet = StringVar(value="")

        options = [
            "Регулярное обновление операционной системы",
            "Использование надёжных паролей",
            "Ограничение времени пребывания в сети",
            "Установка антивирусного ПО"
        ]

        for i, option in enumerate(options):
            Radiobutton(container,
                       text=option,
                       variable=self.otvet,
                       value=option,
                       font=self.radio_font
                       ).grid(row=3+i, column=0, sticky=W, padx=30, pady=5)

        self.textbox2 = Text(container,
                            width=50,
                            height=3,
                            font=self.main_font,
                            wrap=WORD,
                            state='disabled')
        self.textbox2.grid(row=7, column=0, pady=20)

        Button(container,
               text="Ответить",
               command=self.proverka,
               font=self.button_font,
               padx=30,
               pady=5
               ).grid(row=8, column=0, pady=10)


class ItogFrame(Frame):
    def __init__(self, master, root, correct, total):
        super().__init__(master)
        self.root = root
        self.grid()
        self.configure_fonts()
        self.create_widgets(correct, total)

    def configure_fonts(self):
        """Настраивает шрифты для всех элементов"""
        self.main_font = ("Times New Roman", 14)
        self.title_font = ("Times New Roman", 16, "bold")
        self.button_font = ("Times New Roman", 14)

    def create_widgets(self, correct, total):
        """Создает интерфейс для отображения результатов"""
        # Основной контейнер для центрирования
        container = Frame(self)
        container.grid(pady=50)

        percent = round(100 * correct / total)

        # Заголовок
        Label(container,
              text="Тест завершён!",
              font=self.title_font
              ).grid(row=0, column=0, pady=20)

        # Результаты
        Label(container,
              text=f"Правильных ответов: {correct} из {total}",
              font=self.main_font
              ).grid(row=1, column=0, pady=10)

        Label(container,
              text=f"Результат: {percent}%",
              font=self.main_font
              ).grid(row=2, column=0, pady=10)

        # Кнопка выхода
        Button(container,
               text="Выход",
               command=self.quit_app,
               font=self.button_font,
               padx=30,
               pady=5
               ).grid(row=3, column=0, pady=30)

    def quit_app(self):
        """Закрывает приложение"""
        self.root.destroy()


class Final:
    def __init__(self, root):
        self.root = root
        self.question_index = 0
        self.correct_answers = 0

        # Настройка шрифтов для всего приложения
        self.root.option_add("*Font", ("Times New Roman", 14))

        # Список классов вопросов
        self.question_frames = [CheckQuestionFrame, RadioQuestionFrame, TextBoxFrame]
        self.total_questions = len(self.question_frames)

        # Показываем заставку
        self.intro = ZastavkaFrame(self.root, self.start_test)

    def start_test(self):
        """Начинает тест после заставки"""
        self.intro.destroy()
        self.load_question()

    def load_question(self, was_correct=False):
        """Загружает следующий вопрос или завершает тест"""
        # Увеличиваем счетчик правильных ответов
        if self.question_index > 0 and was_correct:
            self.correct_answers += 1

        # Проверяем, есть ли еще вопросы
        if self.question_index < self.total_questions:
            frame_class = self.question_frames[self.question_index]
            self.current_frame = frame_class(self.root, self.load_question, self.question_index)
            self.question_index += 1
        else:
            # Все вопросы пройдены - показываем итоги
            self.save_to_excel()
            self.current_frame = ItogFrame(self.root, self.root, self.correct_answers, self.total_questions)

    def save_to_excel(self):
        """Сохраняет результаты теста в файл Excel"""
        filename = "test_results.xlsx"

        try:
            # Проверяем существование файла
            file_exists = os.path.exists(filename)

            if file_exists:
                wb = load_workbook(filename)
                ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.append(['Дата и время', 'Правильных ответов', 'Всего вопросов', 'Процент'])

            # Добавляем новые данные
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            percent = round(100 * self.correct_answers / self.total_questions)
            ws.append([now, self.correct_answers, self.total_questions, f"{percent}%"])

            # Сохраняем файл
            wb.save(filename)
            print(f"Результаты сохранены в файл {filename}")

        except Exception as e:
            print(f"Ошибка при сохранении в Excel: {e}")

